#!/usr/bin/env python3
"""
Generate items.xlsx (inventory input) and master.xlsx (product BOM + build analysis).

Sources:
  sample_manatec_inventory.csv    -> items.xlsx
  sample_manatec_master_bom.csv   -> master.xlsx

Reuses bom_solver for reading + analysis.
"""

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from bom_solver import analyze_product, read_inventory, read_master

INV_CSV = "sample_manatec_inventory.csv"
MASTER_CSV = "sample_manatec_master_bom.csv"
ITEMS_XLSX = "items.xlsx"
MASTER_XLSX = "master.xlsx"

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(bold=True, color="FFFFFF")
BAD_FILL = PatternFill("solid", fgColor="FFC7CE")
WARN_FILL = PatternFill("solid", fgColor="FFEB9C")
OK_FILL = PatternFill("solid", fgColor="C6EFCE")


def safe_sheet_name(name, used):
    clean = "".join(" " if ch in "[]:*?/\\" else ch for ch in str(name)).strip()
    clean = clean[:31] or "Sheet"
    candidate, n = clean, 1
    while candidate.lower() in used:
        suffix = f"~{n}"
        candidate = clean[: 31 - len(suffix)] + suffix
        n += 1
    used.add(candidate.lower())
    return candidate


def write_table(ws, headers, rows, widths=None):
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row in rows:
        ws.append(row)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for i, header in enumerate(headers, start=1):
        letter = get_column_letter(i)
        ws.column_dimensions[letter].width = (
            widths[i - 1] if widths and i <= len(widths) else max(12, len(header) + 4)
        )


def paint(ws, col_idx, matcher):
    """Apply fill to whole row where matcher(row_values) is True."""
    for row in ws.iter_rows(min_row=2):
        if matcher([c.value for c in row]):
            for cell in row:
                cell.fill = row[col_idx - 1].fill


# ── Load ────────────────────────────────────────────────────────────────────
inventory = read_inventory(INV_CSV, item_col="Item Code", qty_col="Quantity Available")
bom_map = read_master(MASTER_CSV, item_col="Item Code", qty_col="Qty Required Per Unit",
                      product_col="Product")
results = {p: analyze_product(b, inventory) for p, b in bom_map.items()}

# ── items.xlsx ──────────────────────────────────────────────────────────────
wb_items = Workbook()
ws = wb_items.active
ws.title = "Items"
write_table(
    ws,
    ["Item Code", "Quantity Available", "Status"],
    [
        (item, qty, "In Stock" if qty > 0 else "Out of Stock")
        for item, qty in sorted(inventory.items(), key=lambda kv: kv[0].lower())
    ],
    widths=[40, 20, 14],
)
for row in ws.iter_rows(min_row=2):
    if row[1].value == 0:
        for cell in row:
            cell.fill = BAD_FILL
wb_items.save(ITEMS_XLSX)
print(f"wrote {ITEMS_XLSX}  ({len(inventory)} items)")

# ── master.xlsx ─────────────────────────────────────────────────────────────
wb = Workbook()
used = set()

# Sheet 1: the master BOM itself
ws_bom = wb.active
ws_bom.title = safe_sheet_name("Master BOM", used)
write_table(
    ws_bom,
    ["Product", "Item Code", "Qty Required Per Unit"],
    [(p, item, qty) for p in sorted(bom_map) for item, qty in sorted(bom_map[p].items())],
    widths=[42, 40, 22],
)

# Sheet 2: build summary
ws_sum = wb.create_sheet(safe_sheet_name("Build Summary", used))
summary_rows = []
for p in sorted(results):
    r = results[p]
    status = "BUILDABLE" if r["max_units"] > 0 else "BLOCKED"
    summary_rows.append((
        p,
        r["bom_item_count"],
        r["max_units"],
        r["limiting_item"] or "-",
        status,
        len(r["missing_items"]),
        len(r["shortage_details"]),
        len(r["leftover_after_production"]),
    ))
write_table(
    ws_sum,
    ["Product", "BOM Components", "Max Units Buildable", "Limiting Item", "Status",
     "Missing Items", "Short Items", "Leftover Items"],
    summary_rows,
    widths=[42, 18, 20, 32, 14, 16, 14, 16],
)
for row in ws_sum.iter_rows(min_row=2):
    fill = OK_FILL if row[4].value == "BUILDABLE" else BAD_FILL
    for cell in row:
        cell.fill = fill

# Sheet 3: missing / short
ws_miss = wb.create_sheet(safe_sheet_name("Missing Items", used))
missing_rows = []
for p in sorted(results):
    r = results[p]
    for item, need in sorted(r["missing_items"].items()):
        missing_rows.append((p, item, need, inventory.get(item, 0), need, "Not in stock"))
    for item, have, need, short in sorted(r["shortage_details"], key=lambda x: -x[3]):
        missing_rows.append((p, item, need, have, short, "Insufficient stock"))
write_table(
    ws_miss,
    ["Product", "Item Code", "Qty Required Per Unit", "Qty Available", "Qty Short", "Reason"],
    missing_rows or [("-", "-", 0, 0, 0, "None")],
    widths=[42, 40, 22, 16, 12, 22],
)

# Sheet 4: extra / unused inventory
ws_extra = wb.create_sheet(safe_sheet_name("Extra Items", used))
used_items = {i for b in bom_map.values() for i in b}
extra_rows = [
    (item, qty, "Not used in any product")
    for item, qty in sorted(inventory.items(), key=lambda kv: kv[0].lower())
    if item not in used_items
]
write_table(
    ws_extra,
    ["Item Code", "Quantity Available", "Reason"],
    extra_rows or [("-", 0, "None")],
    widths=[40, 20, 30],
)

# Per-product breakdown sheets
for p in sorted(bom_map):
    bom = bom_map[p]
    r = results[p]
    ws = wb.create_sheet(safe_sheet_name(p, used))
    rows = []
    for item, need in sorted(bom.items()):
        have = inventory.get(item, 0)
        rows.append((
            item, need, have,
            "Yes",
            "Yes" if have >= need else "NO",
            have // need if need else 0,
            r["max_units"],
            max(have - r["max_units"] * need, 0),
            max(need - have, 0),
        ))
    write_table(
        ws,
        ["Item Code", "Qty Required/Unit", "Qty Available", "In BOM", "Sufficient",
         "Units From This Item", "Max Units", "Stock After Build", "Shortage"],
        rows,
        widths=[40, 20, 16, 10, 12, 20, 12, 18, 12],
    )
    for row in ws.iter_rows(min_row=2):
        if row[4].value == "NO":
            for cell in row:
                cell.fill = BAD_FILL
        elif row[8].value:
            for cell in row:
                cell.fill = WARN_FILL

wb.save(MASTER_XLSX)
print(f"wrote {MASTER_XLSX}  ({len(bom_map)} products, {len(used)} sheets)")
